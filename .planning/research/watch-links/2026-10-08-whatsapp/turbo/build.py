"""Excel turbo — the 125 WhatsApp links judged by both of the holding's purposes.

One-off builder for this job (the dxb-watch-links door's xlsx.py does not know the revenue
fields; widening the door waits on the CEO's word). Inputs live in the job folder:
  turbo/birlesik.jsonl   reader rows (g1-5) with the lead's corrections already applied
  turbo/split-*.jsonl    the two engineers' tech/revenue split; later lines win per id
  turbo/itiraz-*.jsonl   each engineer's cross-check of the other's half (later file wins)
  turbo/kararlar.jsonl   our recommendation per project and per decision waiting on the CEO
  projeler.jsonl         the nine projects ordered by money
  items.tsv              date, links and the CEO's own note per id
  the second-pass workbook — its Holding/Öneri values are the "before" column, its
  Tür / Tahta satırı / Neye bakıldı cells are carried over.
Usage: python build.py <second-pass.xlsx> <out.xlsx>
"""
import csv
import re
import glob
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

JOB = Path(__file__).resolve().parent.parent
TURBO = JOB / "turbo"

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
SCORE_FILL = {3: "C6EFCE", 2: "FFEB9C", 1: "FCE4D6", 0: "EDEDED"}
ONERI_ORDER = ["Kur", "Dene", "Araştır", "İzle", "Gerek yok"]


def jsonl(path):
    out = []
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line:
            out.append(json.loads(line))
    return out


def as_int(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def load_rows():
    rows = {r["id"]: r for r in jsonl(TURBO / "birlesik.jsonl")}
    for r in rows.values():
        r["gt_puan"], r["gt_oneri"] = as_int(r.get("holding_puan")), r.get("oneri")
    split_files = sorted(glob.glob(str(TURBO / "split-*.jsonl")))
    for f in split_files:
        for s in jsonl(f):
            r = rows[s["id"]]
            r.update({k: s[k] for k in ("tek_katki", "tek_puan", "gelir_katki", "gelir_puan") if k in s})
            for k in ("holding_puan", "oneri"):
                if s.get(k) not in (None, ""):
                    r[k] = s[k]
            if s.get("degisiklik"):
                r["degisiklik_turbo"] = s["degisiklik"]
    # The cross-check verdicts are applied last: each reviewer writes its own itiraz-*.jsonl;
    # when both files carry an id, the file changed later wins.
    for f in sorted(TURBO.glob("itiraz-*.jsonl"), key=lambda f: f.stat().st_mtime):
        for s in jsonl(f):
            r = rows[s["id"]]
            for k in ("tek_puan", "gelir_puan", "holding_puan", "oneri"):
                if s.get(k) not in (None, ""):
                    r[k] = s[k]
            if s.get("karar"):
                r["degisiklik_turbo"] = s["karar"]
            r["iki_gorus"] = s.get("iki_gorus") or ""
    return rows, split_files


def load_before(xlsx):
    wb = load_workbook(xlsx, read_only=True)
    ws = wb["Tüm linkler"]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    col = {h: i for i, h in enumerate(head)}
    before = {}
    for r in it:
        if not r[0]:
            continue
        before[r[0]] = {
            "tur": r[col["Tür"]],
            "puan": as_int(r[col["Holding (0-3)"]]),
            "oneri": r[col["Öneri"]],
            "tahta": r[col["Tahta satırı"]],
            "bakildi": r[col["Neye bakıldı"]],
            "okundu": r[col["Okundu mu"]],
        }
    return before


def load_items():
    with open(JOB / "items.tsv", encoding="utf-8") as f:
        return {r["id"]: r for r in csv.DictReader(f, delimiter="\t")}


def sheet(wb, title, headers, widths, rows, score_cols=(), freeze="D2"):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HDR_FILL, HDR_FONT, WRAP
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
        for ci in score_cols:
            c = row[ci]
            v = as_int(c.value)
            if v is not None:
                c.value = v
                c.fill = PatternFill("solid", fgColor=SCORE_FILL[v])
                c.alignment = Alignment(horizontal="center", vertical="top")
    ws.freeze_panes = freeze
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    return ws


def main(src_xlsx, out_xlsx):
    rows, split_files = load_rows()
    before = load_before(src_xlsx)
    items = load_items()
    projects = sorted(jsonl(JOB / "projeler.jsonl"), key=lambda p: p["sira"])
    ids = sorted(rows)
    missing_split = [i for i in ids if "tek_puan" not in rows[i]]
    # CEO, 2026-10-09: 'Gerek yok' incelendi, gerçekten gerek yok — bu linkler detay sayfalarından
    # tamamen çıkar (Tüm linkler, Değişenler), ayrı "Gerek yok" sayfası kalksın. Yalnız Özet'teki
    # öneri sayımı (süreç kaydı) kalır.
    ids_keep = [i for i in ids if rows[i]["oneri"] != "Gerek yok"]

    wb = Workbook()
    wb.remove(wb.active)

    # --- Tüm linkler -------------------------------------------------------------------
    all_head = [
        "No", "Tarih", "Başlık", "Tür", "Holding (0-3)", "Teknoloji (0-3)", "Teknolojiye katkısı",
        "Gelir (0-3)", "Gelire katkısı", "Gelir katmanı", "Nasıl para getirir (okuyucu)",
        "Rakam kanıtı", "İlk para", "Engel ve aşma", "YouTube izni (A/B)", "İslami sınır", "Öneri",
        "İkinci turda", "Neden değişti", "İki görüş (uzlaşılmadı)", "Lider notu", "Oturum (0-3)", "Tahta satırı",
        "Sizin notunuz", "Link(ler)", "Neye bakıldı",
    ]
    all_w = [6, 11, 34, 11, 9, 9, 50, 9, 50, 22, 60, 45, 45, 45, 35, 30, 11, 14, 45, 40, 45, 9, 14, 35, 40, 40]

    def clean(t):
        # Reviewer shorthand ("Opus t2/g1/h2 → Fable t1/g1/h1.", "Opus 'Araştır' → Fable 'İzle'.") is ours, not his.
        t = re.sub(r"^Opus itirazı:\s*", "", t or "")
        t = re.sub(r"^Opus [^→]*→ Fable [^.]*\.\s*", "", t)
        return t[:1].upper() + t[1:]

    def why(r):
        parts = [clean(r.get("degisiklik_turbo")), r.get("neden_degisti_lider") or r.get("neden_degisti")]
        return " · ".join(p for p in parts if p)

    def full_row(i):
        r, b, it = rows[i], before.get(i, {}), items.get(i, {})
        prev = f"{b.get('puan')}/{b.get('oneri')}" if b else ""
        return [
            i, it.get("date"), r.get("baslik"), b.get("tur"), r.get("holding_puan"), r.get("tek_puan"),
            r.get("tek_katki"), r.get("gelir_puan"), r.get("gelir_katki"), r.get("gelir_katmani"),
            r.get("gelir_yolu"), r.get("rakam_kaniti"), r.get("ilk_para"), r.get("engel_ve_asma"),
            r.get("youtube_izni") or "", r.get("islami_sinir"), r.get("oneri"), prev, why(r),
            r.get("iki_gorus") or "", r.get("lider_notu") or "", r.get("oturum_puan"), b.get("tahta"), it.get("ceo_note") or "",
            it.get("links"), b.get("bakildi") or r.get("okunan"),
        ]

    # --- Özet --------------------------------------------------------------------------
    ws = wb.create_sheet("Özet")
    now = Counter(rows[i]["oneri"] for i in ids)
    was = Counter(before[i]["oneri"] for i in ids if i in before)
    gt = Counter(rows[i]["gt_oneri"] for i in ids)
    tek3 = sum(1 for i in ids if as_int(rows[i].get("tek_puan")) == 3)
    gel3 = sum(1 for i in ids if as_int(rows[i].get("gelir_puan")) == 3)
    both = sum(1 for i in ids if (as_int(rows[i].get("tek_puan")) or 0) >= 2 and (as_int(rows[i].get("gelir_puan")) or 0) >= 2)
    up = sum(1 for i in ids if i in before and (as_int(rows[i]["holding_puan"]) or 0) > (before[i]["puan"] or 0))
    down = sum(1 for i in ids if i in before and (as_int(rows[i]["holding_puan"]) or 0) < (before[i]["puan"] or 0))
    lines = [
        ["Excel turbo — WhatsApp 'Claude Code' grubundaki 125 link, holdingin iki gayesiyle", None],
        ["Gaye: teknolojik gelişmişlik ve gelir üretimi (CEO, 2026-10-08)", None],
        [None, None],
        ["Öneri", "Turbo (şimdi)", "Gelir turu", "İkinci tur"],
        *[[o, now.get(o, 0), gt.get(o, 0), was.get(o, 0)] for o in ONERI_ORDER],
        [None, None],
        ["Holding puanı yükselen (ikinci tura göre)", up],
        ["Holding puanı düşen (ikinci tura göre)", down],
        ["Teknolojiye katkısı güçlü (3)", tek3],
        ["Gelire katkısı güçlü (3)", gel3],
        ["İki tarafı da gerçek (ikisi de ≥ 2)", both],
        [None, None],
        ["Sayfalar", None],
        ["Projeler", "Dokuz gelir projesi, paraya göre sıralı; dayandığı linkler, önerimiz ve sizi bekleyen kararlar"],
        ["Öncelikli", "Holding puanı 2-3 ve öneri Kur / Dene / Araştır olan linkler"],
        ["Tüm linkler", f"{len(ids_keep)} link — '{ONERI_ORDER[-1]}' denen {now.get(ONERI_ORDER[-1], 0)} link çıkarıldı (CEO 2026-10-09, incelendi, gerçekten gerek yok); teknoloji ve gelir katkısı ayrı sütunlarda"],
        ["Değişenler", "Puanı ya da önerisi değişen, 'Gerek yok' ile sonuçlanmayan linkler — üç durak: ikinci tur → gelir turu → turbo, ve nedeni"],
        [None, None],
        ["YouTube sponsorluğu", "Kapsamı bilinmiyor. YouTube'a dayanan satırlarda iki okuma var — A: YouTube'un kendi kuralları · B: başkalarının görüntüsü de serbest. Varsayım yapılmadı."],
    ]
    for l in lines:
        ws.append(l)
    ws["A1"].font = Font(bold=True, size=14)
    for c in ws[4]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # --- Projeler ----------------------------------------------------------------------
    kararlar = jsonl(TURBO / "kararlar.jsonl") if (TURBO / "kararlar.jsonl").exists() else []
    onerim = {str(k["sira"]): k["onerim"] for k in kararlar if k.get("tur") == "proje"}
    ws = sheet(
        wb, "Projeler",
        ["Sıra", "Proje", "Ne yapılır", "İlk para", "Dayandığı linkler", "Tahta satırı", "Sizin onayınız",
         "Risk", "Önerimiz"],
        [6, 32, 70, 55, 30, 20, 40, 45, 55],
        [[p["sira"], p["proje"], p["ne"], p["ilk_para"],
          ", ".join(p["linkler"]) if isinstance(p["linkler"], list) else p["linkler"],
          p["tahta"], p["onay"], p["risk"], onerim.get(str(p["sira"]), "")] for p in projects],
        freeze="C2",
    )
    pending = [k for k in kararlar if k.get("tur") == "karar"]
    if pending:
        ws.auto_filter.ref = None
        ws.append([])
        ws.append([None, "Sizi bekleyen kararlar", "Önerimiz", "Neden"])
        for c in ws[ws.max_row]:
            c.fill, c.font = HDR_FILL, HDR_FONT
        for n, k in enumerate(pending, 1):
            ws.append([n, k["karar"], k["onerim"], k.get("neden", "")])
            for c in ws[ws.max_row]:
                c.alignment = WRAP

    score_cols = (4, 5, 7, 21)
    key = lambda i: (-(as_int(rows[i]["holding_puan"]) or 0),
                     -((as_int(rows[i].get("tek_puan")) or 0) + (as_int(rows[i].get("gelir_puan")) or 0)), i)
    prio = sorted([i for i in ids if (as_int(rows[i]["holding_puan"]) or 0) >= 2
                   and rows[i]["oneri"] in ("Kur", "Dene", "Araştır")], key=key)
    sheet(wb, "Öncelikli", all_head, all_w, [full_row(i) for i in prio], score_cols)
    sheet(wb, "Tüm linkler", all_head, all_w, [full_row(i) for i in ids_keep], score_cols)

    changed = [i for i in ids_keep if i in before and len({
        (before[i]["puan"], before[i]["oneri"]), (rows[i]["gt_puan"], rows[i]["gt_oneri"]),
        (as_int(rows[i]["holding_puan"]), rows[i]["oneri"])}) > 1]
    sheet(
        wb, "Değişenler",
        ["No", "Başlık", "İkinci tur puan", "Gelir turu puan", "Turbo puan", "İkinci tur öneri",
         "Gelir turu öneri", "Turbo öneri", "Teknoloji (0-3)", "Gelir (0-3)", "Neden değişti",
         "İki görüş (uzlaşılmadı)", "Lider notu"],
        [6, 40, 9, 9, 9, 11, 11, 11, 9, 9, 60, 40, 50],
        [[i, rows[i].get("baslik"), before[i]["puan"], rows[i]["gt_puan"], rows[i]["holding_puan"],
          before[i]["oneri"], rows[i]["gt_oneri"], rows[i]["oneri"], rows[i].get("tek_puan"),
          rows[i].get("gelir_puan"), why(rows[i]), rows[i].get("iki_gorus") or "",
          rows[i].get("lider_notu") or ""] for i in changed],
        score_cols=(2, 3, 4, 8, 9), freeze="C2",
    )

    # Print layout: landscape everywhere; the narrow sheets one page wide. The 26-column link
    # sheets are not squeezed — at one page wide their text would be unreadable on paper.
    for ws in wb:
        ws.page_setup.orientation = "landscape"
        if ws.title in ("Özet", "Projeler", "Değişenler"):
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(out_xlsx)
    print(f"rows={len(ids)} split_files={[Path(f).name for f in split_files]} "
          f"missing_split={len(missing_split)} prio={len(prio)} changed={len(changed)} "
          f"now={dict(now)} was={dict(was)} tek3={tek3} gel3={gel3} both={both}")
    if missing_split:
        print("missing split:", ",".join(missing_split))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
