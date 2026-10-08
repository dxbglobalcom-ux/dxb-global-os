# xlsx.py <job-dir> <out.xlsx> — the readers' rows + the lead's plan -> one spreadsheet for the CEO.
# Inputs: items.tsv (id, date, links, ceo_note), rows/*.jsonl (readers; a later file overrides an
# earlier one by id), plan.jsonl (the lead: may correct holding_puan/oturum_puan/oneri/tahta/ne_bu/izlendi, and carries
# plan, nerede, tahta, kim, onay, risk, sira). Needs openpyxl (a venv: python3 -m venv v; v/bin/pip install openpyxl).
import csv, json, glob, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

J, OUT = sys.argv[1], sys.argv[2]
items = {r["id"]: r for r in csv.DictReader(open(f"{J}/items.tsv", encoding="utf-8"), delimiter="\t")}

def jsonl(path):
    for line in open(path, encoding="utf-8"):
        if line.strip():
            try: yield json.loads(line)
            except json.JSONDecodeError: print("bad line in", path, line[:80])

res = {}
for f in sorted(glob.glob(f"{J}/rows/*.jsonl")):
    for o in jsonl(f): res[o["id"]] = o
plan = {o["id"]: o for o in jsonl(f"{J}/plan.jsonl")} if os.path.exists(f"{J}/plan.jsonl") else {}
for i, o in plan.items():
    for k in ("holding_puan", "oturum_puan", "oneri", "tahta", "ne_bu", "izlendi", "okundu"):
        if k in o and i in res: res[i][k] = o[k]
print("items", len(items), "rows", len(res), "plan", len(plan), "missing", [i for i in items if i not in res])

def num(v):
    try: return int(v)
    except Exception: return 0

HEAD = ["No", "Tarih", "Başlık", "Ne bu", "Tür", "Holding (0-3)", "Holding için neden", "Oturum (0-3)",
        "Oturum için neden", "Bizde zaten var mı", "Tahta satırı", "Öneri", "Not", "Uygulama planı",
        "Sizin notunuz", "Link(ler)", "Neye bakıldı", "Okundu mu"]
WIDTH = [6, 11, 32, 55, 13, 9, 45, 9, 45, 35, 11, 11, 45, 60, 40, 45, 30, 25]
FILL = {3: "C6EFCE", 2: "E2F0D9", 1: "FFF2CC", 0: "F2F2F2"}
OFILL = {"Kur": "C6EFCE", "Dene": "DDEBF7", "Araştır": "FFF2CC", "İzle": "F2F2F2", "Gerek yok": "F8CBAD"}

def row(i):
    it, o = items[i], res.get(i, {})
    return [i, it.get("date", ""), o.get("baslik", ""), o.get("ne_bu", ""), o.get("tur", ""),
            num(o.get("holding_puan")), o.get("holding_neden", ""), num(o.get("oturum_puan")),
            o.get("oturum_neden", ""), o.get("zaten_var", ""), o.get("tahta", ""), o.get("oneri", "OKUNMADI"),
            o.get("not", ""), plan.get(i, {}).get("plan", ""), it.get("ceo_note", ""), it.get("links", ""),
            o.get("izlendi", ""), o.get("okundu", "hayır — sonuç yok")]

def head(ws, cols, widths):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for k, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(k)].width = w

def sheet(ws, rows):
    head(ws, HEAD, WIDTH)
    for r in rows: ws.append(r)
    for r in ws.iter_rows(min_row=2):
        for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in (5, 7):
            if isinstance(r[col].value, int): r[col].fill = PatternFill("solid", fgColor=FILL.get(r[col].value, "FFFFFF"))
        if r[11].value in OFILL: r[11].fill = PatternFill("solid", fgColor=OFILL[r[11].value])
    ws.freeze_panes = "D2"; ws.auto_filter.ref = ws.dimensions

wb = Workbook()
ws0 = wb.active; ws0.title = "Özet"
allrows = [row(i) for i in items]
pri = sorted([r for r in allrows if r[11] in ("Kur", "Dene") or r[5] + r[7] >= 4], key=lambda r: (-(r[5] + r[7]), r[0]))
counts = {}
for r in allrows: counts[r[11]] = counts.get(r[11], 0) + 1
ws0.append([f"{len(items)} kayıt; {len(res)} tanesi içeriğinden incelendi; {len(plan)} tanesinin uygulama planı var."])
ws0["A1"].font = Font(bold=True, size=13)
ws0.append([]); ws0.append(["Öneri", "Adet"]); ws0["A3"].font = ws0["B3"].font = Font(bold=True)
for k in ["Kur", "Dene", "Araştır", "İzle", "Gerek yok", "OKUNMADI"]:
    if counts.get(k): ws0.append([k, counts[k]])
ws0.append([]); ws0.append(["Plan = işe yarayanların nasıl uygulanacağı · Öncelikli = Kur/Dene ya da toplam puanı 4+ · Tüm linkler = tarih sırasıyla hepsi."])
ws0.column_dimensions["A"].width = 30; ws0.column_dimensions["B"].width = 10

wp = wb.create_sheet("Plan")
head(wp, ["Sıra", "No", "Başlık", "Ne yapılacak", "Nerede", "Tahta satırı", "Kim yapar", "Onay", "Maliyet ve risk"],
     [6, 6, 32, 60, 35, 12, 22, 35, 40])
for o in sorted([o for o in plan.values() if o.get("plan")], key=lambda o: (o.get("sira", 99), o["id"])):
    wp.append([o.get("sira", ""), o["id"], res.get(o["id"], {}).get("baslik", ""), o.get("plan", ""), o.get("nerede", ""),
               o.get("tahta", ""), o.get("kim", ""), o.get("onay", ""), o.get("risk", "")])
for r in wp.iter_rows(min_row=2):
    for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")
wp.freeze_panes = "D2"
sheet(wb.create_sheet("Öncelikli"), pri)
sheet(wb.create_sheet("Tüm linkler"), allrows)
wb.save(OUT)
print("saved", OUT, "priority", len(pri), counts)
